"""
Generates Dermify_Project_Overview.pptx — a capstone-defence deck covering the
app's architecture and the data curation pipeline.

Content is sourced only from DERMIFY_PROJECT_CONTEXT.md, DATASET_CURATION_
PIPELINE.md, and facts verified during the curation-pipeline debugging session
(the real Postgres/Supabase bugs found and fixed). Nothing here is invented.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------------------
# Palette — teal matches the app's own colorPrimary (#0097A7)
# ---------------------------------------------------------------------------
TEAL = RGBColor(0x00, 0x97, 0xA7)
TEAL_DARK = RGBColor(0x00, 0x6B, 0x78)
INK = RGBColor(0x1C, 0x26, 0x28)
SLATE = RGBColor(0x4A, 0x58, 0x5B)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
PAPER_TINT = RGBColor(0xF2, 0xFA, 0xFB)
ACCENT = RGBColor(0xE0, 0x5A, 0x00)
LINE = RGBColor(0xD8, 0xE6, 0xE8)

FONT = "Calibri"
FONT_LIGHT = "Calibri Light"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_footer(slide, page_no):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SLIDE_H - Inches(0.28), SLIDE_W, Inches(0.28))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = f"Dermify — On-Device Skin Screening & Data Curation Pipeline        {page_no:02d}"
    p.font.size = Pt(10)
    p.font.color.rgb = PAPER
    p.font.name = FONT
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)


def add_accent_bar(slide, y=Inches(1.15)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), y, Inches(1.1), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title_text(slide, text, top=Inches(0.45), size=32, color=INK, bold=True, left=Inches(0.5), width=Inches(12.33)):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT
    return box


def add_kicker(slide, text, top=Inches(0.15)):
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(10), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text.upper()
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = FONT
    run = p.runs[0]
    run.font._rPr.set('spc', '150')


def add_bullets(slide, items, top=Inches(1.55), left=Inches(0.5), width=Inches(12.3), height=Inches(5.3),
                 size=20, line_spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            level, text = item
        else:
            level, text = 0, item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.line_spacing = line_spacing
        p.space_after = Pt(10 if level == 0 else 4)
        bullet_char = "●" if level == 0 else "–"
        p.text = f"{bullet_char}  {text}"
        p.font.size = Pt(size if level == 0 else size - 3)
        p.font.color.rgb = INK if level == 0 else SLATE
        p.font.name = FONT
        p.font.bold = (level == 0)
    return box


def add_callout(slide, text, top=Inches(6.15), left=Inches(0.5), width=Inches(12.3), height=Inches(0.85)):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = PAPER_TINT
    box.line.color.rgb = TEAL
    box.line.width = Pt(1)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(16)
    p.font.italic = True
    p.font.bold = True
    p.font.color.rgb = TEAL_DARK
    p.font.name = FONT
    p.alignment = PP_ALIGN.LEFT
    return box


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_content_slide(page_no, kicker, title, notes):
    s = add_slide()
    set_bg(s)
    add_kicker(s, kicker)
    add_title_text(s, title)
    add_accent_bar(s)
    add_footer(s, page_no)
    add_notes(s, notes)
    return s


def add_table(slide, headers, rows, top=Inches(1.7), left=Inches(0.6), width=Inches(12.1), height=Inches(3.6),
              col_widths=None, highlight_rows=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gtable = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = gtable.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = PAPER
        p.font.name = FONT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    highlight_rows = highlight_rows or set()
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PAPER_TINT if r in highlight_rows else PAPER
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(14)
            p.font.name = FONT
            p.font.color.rgb = INK
            p.font.bold = (r in highlight_rows)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return gtable


def add_flow(slide, steps, top=Inches(2.6), box_h=Inches(1.3), gap=Inches(0.15)):
    n = len(steps)
    total_w = Inches(12.3)
    box_w = Emu(int((total_w - gap * (n - 1)) / n))
    x = Inches(0.5)
    for i, step in enumerate(steps):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = TEAL if i % 2 == 0 else TEAL_DARK
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = step
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PAPER
        p.font.name = FONT
        p.alignment = PP_ALIGN.CENTER
        if i < n - 1:
            arrow = slide.shapes.add_textbox(Emu(int(x + box_w)), Emu(int(top + box_h / 2 - Inches(0.2))), gap, Inches(0.4))
            ap = arrow.text_frame.paragraphs[0]
            ap.text = "→"
            ap.font.size = Pt(20)
            ap.font.bold = True
            ap.font.color.rgb = ACCENT
            ap.alignment = PP_ALIGN.CENTER
        x = Emu(int(x + box_w + gap))


# ===========================================================================
# 1. TITLE
# ===========================================================================
s = add_slide()
set_bg(s, TEAL)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.55), SLIDE_W, Inches(2.95))
band.fill.solid()
band.fill.fore_color.rgb = PAPER
band.line.fill.background()

bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(2.15), Inches(1.4), Pt(6))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

t = s.shapes.add_textbox(Inches(0.65), Inches(2.35), Inches(11.5), Inches(1.4))
p = t.text_frame.paragraphs[0]
p.text = "Dermify"
p.font.size = Pt(60)
p.font.bold = True
p.font.color.rgb = PAPER
p.font.name = FONT_LIGHT

t2 = s.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(11.7), Inches(1.0))
p2 = t2.text_frame.paragraphs[0]
p2.text = "On-device AI skin-condition screening, and a consented pipeline\nto build the dataset it actually needs"
p2.font.size = Pt(20)
p2.font.color.rgb = PAPER
p2.font.name = FONT
t2.text_frame.word_wrap = True

t3 = s.shapes.add_textbox(Inches(0.7), Inches(5.1), Inches(11), Inches(1.6))
tf3 = t3.text_frame
tf3.word_wrap = True
lines = [
    ("BSc Telecommunication Engineering — Final-Year Capstone", 18, True, INK),
    ("Kwame Nkrumah University of Science and Technology (KNUST), Kumasi, Ghana", 15, False, SLATE),
    ("Caleb — Android / Kotlin        Hannah — Python / Machine Learning", 15, False, SLATE),
]
for i, (txt, sz, bold, col) in enumerate(lines):
    p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
    p.text = txt
    p.font.size = Pt(sz)
    p.font.bold = bold
    p.font.color.rgb = col
    p.font.name = FONT
    p.space_after = Pt(8)

add_notes(s, "Opening: Dermify is an Android app that screens photographs of skin for four "
             "conditions, runs entirely on-device, explains its prediction with a heatmap, and "
             "refers uncertain cases to a clinician over WhatsApp. Today's talk covers two things: "
             "how the app itself works end to end, and the data curation pipeline built to fix its "
             "biggest limitation — training data.")

# ===========================================================================
# 2. THE PROBLEM
# ===========================================================================
s = new_content_slide(2, "Motivation", "Why This Exists",
    "Ground the project in a real access problem, not a generic 'AI for healthcare' pitch. "
    "Dermatologist access in Ghana is limited relative to skin-condition burden. The app is "
    "designed as a screening-and-referral tool, not a replacement for a clinician — low-confidence "
    "results are explicitly routed to a specialist rather than presented as a diagnosis.")
add_bullets(s, [
    "Skin conditions are common in Ghana; specialist dermatologist access is limited",
    "Dermify screens a photo on-device and gives an immediate, explainable result",
    "Low-confidence or uncertain cases are referred to a clinician over WhatsApp — the app "
    "augments access, it does not replace a dermatologist",
    "The four conditions it screens for were chosen from real Ghanaian hospital prevalence "
    "data, not from whatever public dataset was convenient",
], size=21)
add_callout(s, "Design principle: built around what Ghanaian patients actually present with, "
               "not around what Western datasets happen to contain.")

# ===========================================================================
# 3. THE FOUR CONDITIONS (table)
# ===========================================================================
s = new_content_slide(3, "Clinical Basis", "Four Conditions, Chosen From Prevalence Data",
    "Sources: KBTH Dermatology Clinic review (Rosenbaum et al., 2017); KATH study "
    "(Asiedu-Gyekye et al., 2022, n=460); community study (Amoako et al., BMC Infect Dis 2024, "
    "n=1,327). Point out the inverse relationship deliberately — eczema and fungal infection "
    "are the two highest-prevalence conditions AND the two lowest F1 scores. This single fact "
    "motivates the entire second half of the talk — the data curation pipeline.")
add_table(s,
    ["Condition", "Prevalence in Ghana", "Model F1 (5-fold CV)"],
    [
        ["Eczema / dermatitis", "24.6%", "69.4%"],
        ["Fungal infection", "15.2%", "71.4%"],
        ["Scabies", "5.1%", "79.3%"],
        ["Acne vulgaris", "5.3%", "92.7%"],
    ],
    top=Inches(1.75), height=Inches(2.6),
    col_widths=[Inches(5.5), Inches(3.3), Inches(3.3)],
    highlight_rows={0, 1},
)
add_callout(s, "The two conditions Ghana most needs classified are the two the model is worst "
               "at — the motivation for the curation pipeline covered later.",
            top=Inches(4.7))

# ===========================================================================
# 4. SYSTEM AT A GLANCE
# ===========================================================================
s = new_content_slide(4, "Architecture", "System at a Glance",
    "Walk left to right. Capture goes through a three-layer gate before it's even allowed to "
    "reach the disease model — this rejects non-skin photos cheaply before spending a full "
    "EfficientNet forward pass. The disease model produces both a class prediction and a "
    "heatmap in one pass. From the result screen, two independent paths exist: refer to a "
    "specialist (telemedicine), or consent to contribute the image for future training — "
    "entirely separate decisions, one does not require the other.")
add_flow(s, ["Capture\n(camera / gallery)", "3-Layer\nCapture Gate", "On-Device\nDisease Model",
             "Prediction +\nCAM Heatmap", "Result Screen"], top=Inches(1.9), box_h=Inches(1.15))
add_flow(s, ["Low confidence /\nuser request", "Telemedicine Referral\n(WhatsApp, de-identified)"],
         top=Inches(3.55), box_h=Inches(1.0))
add_flow(s, ["Explicit consent\n(opt-in)", "De-identify +\nHash (on device)", "Curation Pipeline\n→ Supabase"],
         top=Inches(4.85), box_h=Inches(1.0))
add_callout(s, "Referral and contribution are independent, opt-in paths off the same result screen "
               "— neither is required to get your result.", top=Inches(6.1), height=Inches(0.75))

# ===========================================================================
# 5. TECH STACK
# ===========================================================================
s = new_content_slide(5, "Implementation", "Android App Stack",
    "Keep this brief — it's a reference slide, not a talking point to dwell on, unless asked. "
    "The single-TFLite-dependency constraint and the Room build flags are both things that broke "
    "the build before being pinned down — worth mentioning briefly if asked about engineering "
    "challenges.")
add_bullets(s, [
    "Kotlin, Material 3 (DayNight theme), single navigation graph",
    "ML runtime: TensorFlow Lite 2.17.0 — kept to a single ML dependency to avoid AGP 9.1 "
    "namespace conflicts",
    "Local storage: Room (scan history + the curation outbox queue)",
    "Identity: Firebase Auth + Firestore, three-tier roles (Clinician / Specialist / Admin)",
    "Cloud storage: Supabase (Storage + Postgrest), West Europe region",
    "Target device: Samsung SM-A320FL, Android 8 — the low-end phone that sets the latency budget",
], size=20)

# ===========================================================================
# 6. CAPTURE GATE
# ===========================================================================
s = new_content_slide(6, "On-Device Pipeline", "Is This Even Skin? — The Capture Gate",
    "All three layers must pass (AND), cheapest first. Each is wrapped in try/catch — if a "
    "model fails to load, that layer is skipped rather than blocking the user, since earlier "
    "layers already passed. Worth mentioning: Layer 3's behaviour (which direction the sigmoid "
    "output actually meant) was determined empirically because the documented description of the "
    "model turned out to be inverted — a good example of 'verify, don't trust the description.'")
add_bullets(s, [
    "Layer 1 — SkinHeuristic: YCbCr colour-ratio maths, ~50ms, no model. Catches rooms, sky, "
    "grass, documents",
    "Layer 2 — NotSkinGate: MobileNet v1 (ImageNet). A confident ImageNet-object prediction "
    "means it's a recognisable object, not skin",
    "Layer 3 — SkinBinaryGate: MobileNetV2, domain-trained sigmoid. Threshold and output "
    "direction were determined empirically, not from documentation — the written description "
    "was inverted from what the model actually did",
    "Models are loaded once via ModelCache and reused — constructing a fresh interpreter per "
    "scan was a real, measured latency cost before that fix",
], size=19)

# ===========================================================================
# 7. DISEASE MODEL
# ===========================================================================
s = new_content_slide(7, "Machine Learning", "The Disease Classification Model",
    "EfficientNetB0 with ImageNet weights, custom head. Two-phase training is standard "
    "transfer-learning practice but worth explaining why BatchNorm stays frozen during "
    "fine-tuning: unfreezing BN layers with a small batch destabilises the running statistics. "
    "The dual-output export is a genuine engineering decision — it means the on-device heatmap "
    "costs zero extra inference passes.")
add_bullets(s, [
    "EfficientNetB0 (ImageNet weights) → GlobalAveragePooling → BatchNorm → "
    "Dense(256)+Dropout(0.5) → Dense(128)+Dropout(0.4) → Dense(4, softmax)",
    "Input 224×224. Loss: categorical crossentropy, label smoothing 0.05",
    (1, "Phase 1 — base frozen, Adam 1e-3, 30 epochs"),
    (1, "Phase 2 — unfrozen from layer 165 (BatchNorm layers stay frozen), Adam 2e-5 with "
        "cosine decay, 30 epochs"),
    "Exported as a dual-output TFLite model: class logits AND 7×7×1280 conv features, in one "
    "forward pass — this is what makes the on-device heatmap free",
], size=19)

# ===========================================================================
# 8. EXPLAINABILITY
# ===========================================================================
s = new_content_slide(8, "Explainability", "CAM, Not Grad-CAM — A Deliberate Distinction",
    "This is a good technical-depth slide for a defence — it shows an understanding of a subtle "
    "but important distinction, not just 'we added a heatmap.' TFLite genuinely cannot run "
    "backpropagation at inference time, so a compromise was engineered: precompute class-averaged "
    "gradient weights once in the notebook (using 60 training images per class), then apply those "
    "fixed weights to any new image's feature maps at inference. It's coarser than true Grad-CAM "
    "but works within TFLite's constraints.")
add_bullets(s, [
    "In the training notebook: true Grad-CAM — per-image gradients via GradientTape",
    "In the deployed app: TFLite cannot compute gradients at inference, so the notebook "
    "exports class-averaged pooled gradients (60 images/class) as fixed weights",
    "The app multiplies its exported feature maps by those fixed per-class weights — no "
    "backprop needed on-device",
    "Trade-off, stated plainly: the app's heatmap shows \"what generally matters for this "
    "class,\" not \"what mattered for this specific image\" — coarser, and can miss "
    "atypically-located lesions",
], size=19)

# ===========================================================================
# 9. RESULTS
# ===========================================================================
s = new_content_slide(9, "Evaluation", "Results — 78.14% ± 0.83",
    "Lead with the headline number and its uncertainty, not a single best-run figure — this is "
    "itself a defensible methodological choice worth stating explicitly: a single held-out split "
    "gave 81.14%, but varied by over a point across runs, so the cross-validated figure is the "
    "one reported, precisely because it can be defended under questioning.")
add_bullets(s, [
    "Group-aware 5-fold cross-validation with test-time augmentation: 78.14% ± 0.83 "
    "(95% CI 77.4–78.9%). Without TTA: 76.83% ± 0.82",
    "Mean AUC-ROC: 0.9571",
    "Dataset: 1,976 images → 1,915 unique groups after perceptual-hash deduplication "
    "(61 redundant copies removed)",
    "Tested and found NOT to help: EfficientNetB1 at 240px (worse, 75.20% ± 1.68); stronger "
    "regularisation (closed the train/val gap but not accuracy)",
    "Conclusion: three materially different configurations all landed in the 74–78% band — "
    "the ceiling is dataset size and label quality, not model capacity",
], size=18)
add_callout(s, "The app ships single-view inference (no TTA) — TTA would roughly 4× latency, "
               "which the Samsung A320 can't absorb.", top=Inches(6.3), height=Inches(0.7))

# ===========================================================================
# 10. LEAKAGE INCIDENT
# ===========================================================================
s = new_content_slide(10, "Methodological Rigor", "A Lesson in Verifying, Not Assuming",
    "This is one of the strongest slides for a defence — it demonstrates real scientific "
    "process: noticing an anomaly, investigating rather than dismissing it, finding a real "
    "(if secondary) problem, fixing it properly, and then discovering the anomaly's actual "
    "cause was something else entirely. Be ready to explain both findings clearly if pressed.")
add_bullets(s, [
    "Anomaly noticed: test accuracy (83.5%) was above validation accuracy (79.4%) — "
    "backwards, and correctly treated as suspicious rather than good news",
    "Investigation found 30 cross-dataset near-duplicate images — the same clinical photo "
    "appearing in different public datasets under different filenames",
    "Fixed properly: perceptual-hash clustering (DCT-based, Hamming ≤5) + union-find, then "
    "split by group so every copy of an image stays in one split",
    "The actual cause of the anomaly turned out to be different: comparing training-time "
    "val_accuracy against a differently-computed test accuracy — fixed by scoring all three "
    "splits through an identical code path",
], size=18)
add_callout(s, "That same perceptual-hash algorithm is reused, unchanged, in the live "
               "collection pipeline's duplicate rejection — covered next.",
            top=Inches(6.3), height=Inches(0.7))

# ===========================================================================
# 11. BEYOND CLASSIFICATION
# ===========================================================================
s = new_content_slide(11, "Product Features", "Beyond Classification",
    "Quick tour of the features that make this a usable clinical tool rather than a pure "
    "research demo.")
add_bullets(s, [
    "Confidence gate — predictions below 50% show a caution note, tying directly into the "
    "telemedicine referral path",
    "Telemedicine referral — de-identified composite image (photo + prediction + heatmap, "
    "no name/patient ID), sent via WhatsApp's native share; store-and-forward, works over "
    "intermittent connectivity",
    "Explicit consent flow, entirely separate from the diagnosis flow — a user gets their "
    "result whether or not they contribute an image",
    "Firebase Auth + Firestore: three-tier roles — Clinician, Specialist, Admin",
    "Room-backed local scan history; ModelCache releases TFLite interpreters under memory "
    "pressure",
], size=19)

# ===========================================================================
# 12. WHY A CURATION PIPELINE
# ===========================================================================
s = new_content_slide(12, "Part Two — Data Curation Pipeline", "The Data Problem",
    "Transition slide into part two of the talk. Reconnect explicitly to slide 3's table — "
    "this isn't a new problem, it's the same one, and the curation pipeline is the direct "
    "response to it.")
add_bullets(s, [
    "Accuracy ceiling is data, not architecture — shown on slide 9: three different model "
    "configurations all landed in the same 74–78% band",
    "Training data comes from public, Western-skewed dermatology datasets — roughly 500 "
    "images per class, with no dermatologist verification",
    "The two conditions Ghana most needs classified are the two the model is worst at",
], size=22)
add_callout(s, "The pipeline's purpose: turn real app usage into a consented, de-identified, "
               "expert-labelled, locally-representative dataset.", top=Inches(4.6), height=Inches(1.0))

# ===========================================================================
# 13. ON-DEVICE PIPELINE
# ===========================================================================
s = new_content_slide(13, "Curation Pipeline", "On-Device: From Consent to Queued Upload",
    "Walk through each stage. The label at this stage is the one the app user already "
    "confirmed or corrected on the Verify screen — not a raw, unreviewed model guess. "
    "De-identification runs before upload with no exceptions; it's written as assertions "
    "that fail the contribution outright, since privacy failures are silent by default — "
    "an EXIF field that survives produces no error, so the code checks explicitly rather "
    "than trusting the strip to have worked.")
add_flow(s, ["Verify screen\n(confirm/correct label)", "Consent\n(opt-in, versioned)",
             "De-identify\n(EXIF/GPS strip,\ndownscale, face reject)", "Perceptual hash\n(on-device DCT)",
             "Room outbox\n(survives restart)"], top=Inches(2.0), box_h=Inches(1.5))
add_bullets(s, [
    "WorkManager drains the outbox — WiFi-preferred by default, exponential backoff, "
    "survives an app kill or a dead connection",
], top=Inches(4.0), size=18)
add_callout(s, "Consent and referral are independent, opt-in decisions — a user gets their "
               "result regardless of either choice.", top=Inches(6.15), height=Inches(0.75))

# ===========================================================================
# 14. SERVER-SIDE PIPELINE
# ===========================================================================
s = new_content_slide(14, "Curation Pipeline", "Server-Side: Supabase",
    "Two RPCs, in this order and not the other. register_contribution does the duplicate check "
    "BEFORE any bytes transfer — meaningful on a metered connection, since a duplicate costs one "
    "small request instead of a full image upload. mark_contribution_stored is a genuine "
    "verification step, not a rubber stamp — it checks the object actually exists in storage "
    "before it will mark a row as stored, which matters and will come up again on the debugging "
    "slide.")
add_flow(s, ["register_contribution\n(global pHash dedup\ncheck)", "Storage upload\n(reserved path)",
             "mark_contribution_stored\n(verifies bytes\nactually exist)"],
         top=Inches(2.0), box_h=Inches(1.4))
add_bullets(s, [
    f"Images organised by the verified label: contributions/<label>/<uuid>.jpg",
    "No table read access for contributors, enforced by the database (Row-Level Security + "
    "SECURITY DEFINER functions) — not just by the app choosing not to ask",
], top=Inches(3.9), size=18)

# ===========================================================================
# 15. PRIVACY BY DESIGN
# ===========================================================================
s = new_content_slide(15, "Privacy & Ethics", "Privacy By Design, Not By Policy",
    "Emphasise that these are enforced, not merely promised — the anon key the app ships with "
    "has no read access to the collection at all, only insert; that's a database guarantee, "
    "not client behaviour a user has to trust. Withdrawal is honest about its limit rather than "
    "overpromising — worth stating plainly if asked about GDPR-style 'right to be forgotten.'")
add_bullets(s, [
    "Never leaves the device unmodified: raw camera file, EXIF, GPS, timestamps finer than a "
    "date, Firebase UID, device identifier",
    "Face-containing images are rejected outright — checked again at contribution time, as a "
    "backstop independent of the capture-time face toggle",
    "Consent is versioned (e.g. consent-v1-2026-08) and recorded locally before any upload is "
    "attempted, so the record survives an upload failure",
    "Withdrawal deletes the stored image immediately — and the consent text says plainly that "
    "an image already baked into a published, trained model cannot be un-trained",
    "Collection at scale requires institutional ethical clearance; until granted, "
    "team-members-only, stated openly rather than quietly collecting",
], size=17)

# ===========================================================================
# 16. DEBUGGING STORY
# ===========================================================================
s = new_content_slide(16, "Engineering Process", "Finding the Real Bug — Three Layers Deep",
    "This is the strongest 'we actually engineered this, we didn't just assemble a tutorial' "
    "slide in the deck — use it if there's time or if pressed on process. Each of these three "
    "was found by measuring against the live system, not by reasoning from the spec — and "
    "notably, the FIRST hypothesis (a missing INSERT policy — the most-likely candidate written "
    "into the original spec) turned out to be wrong. That's worth saying explicitly: the "
    "original guess was disproven by evidence, not confirmed by it.")
add_bullets(s, [
    "1 — A leftover dashboard policy let any contributor read the entire bucket, not just "
    "their own upload. Found by inspecting live policies directly, not assumed. Dropped.",
    "2 — Uploads failed with an RLS error even though the write policy plainly matched. Root "
    "cause: Postgres implements \"upsert\" as INSERT…ON CONFLICT DO UPDATE, which needs a base "
    "SELECT privilege the app's role didn't have — separate from, and in addition to, any RLS "
    "policy. Fixed by removing upsert; it wasn't needed anyway",
    "3 — A false positive nearly shipped: rows were marked \"stored\" with zero actual bytes "
    "anywhere in the bucket. Fixed by making the server verify the object exists before it will "
    "trust the client's claim of success",
], size=17)

# ===========================================================================
# 17. MULTI-DEVICE ARCHITECTURE
# ===========================================================================
s = new_content_slide(17, "Scale", "One Shared Dataset, Many Contributors",
    "If asked 'did you test this with multiple physical devices' — be precise: this is an "
    "architectural guarantee (the Supabase URL and key are constants compiled into the APK, "
    "identical on every install, no per-device code path exists), verified thoroughly on one "
    "device, not something that needed two phones side-by-side to confirm.")
add_bullets(s, [
    "Every install of the app shares the same Supabase project — not per-device or per-login "
    "configuration",
    "No contribution carries any device or user identity — even the team running the project "
    "cannot tell which phone or which account sent a given image",
    "Deduplication is global across every contributor, not per-device — the same near-duplicate "
    "problem that once contaminated the training/test split (slide 10) is now caught at "
    "collection time",
], size=21)

# ===========================================================================
# 18. FEDERATED LEARNING — REJECTED
# ===========================================================================
s = new_content_slide(18, "Alternatives Considered", "Federated Learning — Assessed and Rejected",
    "Include this specifically to preempt a common panel question ('why not federated "
    "learning, since privacy matters here'). The answer demonstrates the team understood what "
    "FL actually solves and correctly identified that it doesn't address this project's actual "
    "bottleneck.")
add_bullets(s, [
    "FL exists so training data never leaves the device — but this project's problem is the "
    "opposite: there is no local labelled data at all, only unlabelled photos",
    "The real bottleneck is expert labelling, which is a human constraint FL doesn't touch",
    "A Samsung A320 on Android 8 cannot realistically run backpropagation through "
    "EfficientNetB0 — thermal, battery and memory limits hit before accuracy would improve",
    "FL needs hundreds to thousands of participating clients to average out gradient noise; "
    "a capstone deployment has a handful",
], size=19)
add_callout(s, "Chosen instead: consented, centralised collection with on-device "
               "de-identification — achieves the relevant privacy properties without FL's "
               "coordination cost.", top=Inches(6.15), height=Inches(0.85))

# ===========================================================================
# 19. STATUS & FUTURE WORK
# ===========================================================================
s = new_content_slide(19, "Status", "What's Built, What's Next",
    "Be precise about status — this is a capstone defence, overclaiming an unbuilt feature is "
    "a bigger risk than admitting a phase isn't done yet. Phase 1 is genuinely built and was "
    "confirmed end-to-end on a physical device this session: a real image was captured, "
    "de-identified, hashed, uploaded, organised by label, and independently confirmed present "
    "in storage by direct database query.")
add_bullets(s, [
    "Phase 1 — built and confirmed end-to-end on device: consent, de-identification, "
    "on-device perceptual hashing, Room outbox + WorkManager upload, server-side global "
    "deduplication, database-enforced privacy lockdown",
    "Phase 2 (next) — dermatologist double-blind labelling screen, two-independent-labels "
    "agreement logic, Cohen's κ inter-rater agreement tracking",
    "Phase 3 — immutable dataset version snapshots, withdrawal propagation, class-aware "
    "contribution prioritisation",
    "Phase 4 — retraining stays deliberately manual and human-gated: a new model must beat the "
    "incumbent's cross-validated accuracy by more than one standard deviation before it ships",
], size=18)

# ===========================================================================
# 20. THANK YOU
# ===========================================================================
s = add_slide()
set_bg(s, TEAL)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(3.0), Inches(1.4), Pt(6))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
t = s.shapes.add_textbox(Inches(0.65), Inches(3.2), Inches(11.5), Inches(1.2))
p = t.text_frame.paragraphs[0]
p.text = "Thank You"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = PAPER
p.font.name = FONT_LIGHT
t2 = s.shapes.add_textbox(Inches(0.7), Inches(4.35), Inches(11.5), Inches(1.0))
p2 = t2.text_frame.paragraphs[0]
p2.text = "Questions"
p2.font.size = Pt(22)
p2.font.color.rgb = PAPER
p2.font.name = FONT
add_notes(s, "Close and open the floor. Likely questions to be ready for: multi-device "
             "architecture (slide 17), why not federated learning (slide 18), how retraining "
             "is gated (slide 19), and the debugging story (slide 16) if process/rigor is asked "
             "about directly.")

prs.save("Dermify_Project_Overview.pptx")
print(f"Saved Dermify_Project_Overview.pptx — {len(prs.slides._sldIdLst)} slides")
